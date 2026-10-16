import openpyxl
from openpyxl.styles import Alignment, Border, Side, Font
from openpyxl.utils import get_column_letter as L
FORM = {1:('FA1','FA07'),2:('FA2','FaA18'),5:('EH','EH04'),6:('BS','BS23'),7:('MB','MB09'),8:('FA1','FA08'),9:('FA2','FA19'),12:('EH','FF'),13:('BS','BS27'),14:('MB','UB02'),
15:('FA1','FA09'),16:('FA2','FA12'),19:('EH','FR01'),20:('BS','BS21'),21:('MB','UB05'),22:('FA1','FA10'),23:('FA2','FA13'),26:('EH','EH01'),27:('BS','BS23'),28:('MB','MB01'),29:('FA1','FA02'),30:('FA2','FA14')}
NAMES = ['สมชาย ใจดี','วิไล รักงาน','ประเสริฐ ตั้งใจ','มานะ อดทน','สุดา ขยัน','กิตติ มั่นคง','อรุณี แสงทอง','ธนา พากเพียร','ปรีชา ฉลาด','นภา สดใส','วีระ กล้าหาญ',
         'จันทร์เพ็ญ ศรีสุข','สมศักดิ์ ทองดี','รัตนา บุญมา','ชัยวัฒน์ สุขใจ','พิมพ์ชนก แก้วใส','อนุชา ใจกว้าง','ศิริพร เพชรดี','ณัฐพล รุ่งเรือง','กมลวรรณ สายบุญ','ธีรวัฒน์ ชื่นชม','วรรณา ดีมาก']
def build(path, mutate=None):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'BS Oct'
    thin = Side(style='thin'); B = Border(left=thin,right=thin,top=thin,bottom=thin)
    ws['A1'] = 'แผนงานของ Body Shop ประจำเดือน  .......Oct..... ปี ...2026'; ws.merge_cells('A1:AF1'); ws['A1'].font = Font(size=16, bold=True); ws['A1'].alignment = Alignment(horizontal='center')
    ws['A2'] = 'คำแนะนำในการกรอกข้อมูล:'; ws['A3'] = '1. วางแผนตามแผนการผลิตรายเดือน'
    ws['A4'] = '2. ลำดับการจัดตาราง: เริ่มจากจุดที่มีปัญหาในเดือนก่อน → จุดระดับ A → จุดระดับ B/C'
    ws['A5'] = '3. ให้จัดตารางจุดที่มีปัญหาและจุดระดับ A วันละ 1 จุด ส่วนจุดระดับ B/C ต้องจัดให้ครบทุกตำแหน่งภายใน 2 เดือน'
    ws['A6'] = 'เดือน    Oct       ปี 2026'; ws.merge_cells('A6:AF6'); ws['A6'].alignment = Alignment(horizontal='center')
    labels = {7:'วันที่',8:'ทีม',9:'ชื่อพนักงาน',11:'หมายเลขตำแหน่งงาน',12:'ระดับตำแหน่งงาน',13:'จำนวนปัญหาที่เก็บได้',14:'หมายเหตุ'}
    for r,t in labels.items(): ws.cell(r,1,t)
    ws.merge_cells('A9:A10')   # tall name row made of two rows
    ws.column_dimensions['A'].width = 20
    rows = {}
    for i,d in enumerate(sorted(FORM)):
        rows[d] = dict(team=FORM[d][0], name=NAMES[i], station=FORM[d][1], level='A' if i%3==0 else ('B' if i%3==1 else 'C'))
    if mutate: mutate(rows)
    for d in range(1,32):
        c = d+1
        ws.cell(7,c,d); ws.column_dimensions[L(c)].width = 6.5
        r = rows.get(d)
        ws.merge_cells(start_row=9,start_column=c,end_row=10,end_column=c)
        if r:
            ws.cell(8,c,r['team']); ws.cell(9,c,r['name']); ws.cell(11,c,r['station']); ws.cell(12,c,r['level'])
        ws.cell(9,c).alignment = Alignment(text_rotation=90, wrap_text=True, vertical='center', horizontal='center')
        ws.cell(11,c).alignment = Alignment(text_rotation=90, vertical='center', horizontal='center')
        for rr in range(7,15): ws.cell(rr,c).border = B
    for rr in range(7,15): ws.cell(rr,1).border = B
    ws.row_dimensions[9].height = 60; ws.row_dimensions[10].height = 60; ws.row_dimensions[11].height = 60; ws.row_dimensions[14].height = 50
    ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(path)
build('/tmp/claude-0/fx/BS_Oct2026.xlsx')
def bad(rows):
    rows[8]['station'] = ''; rows[14]['name'] = ''; del rows[21]
build('/tmp/claude-0/fx/BS_Oct2026_incomplete.xlsx', bad)
print('ok')
