import base64
import io
import logging
import os
from pathlib import Path

from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, render
from django.template.loader import render_to_string
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from org import permissions as perm

from . import services as S
from .models import Audit, AuditStatus, Issue, IssueStatus, IssueType

log = logging.getLogger(__name__)
SYM = {"pass": "√", "fail": "×", "na": "N/A", "": ""}
FONT_DIR = Path(__file__).resolve().parent / "static" / "pjo" / "fonts"


@login_required
def reports(request):
    wss = perm.visible_workshops(request.user)
    code = request.GET.get("workshop", "")
    scope = wss.filter(code=code) if code else wss
    audits = Audit.objects.filter(workshop__in=scope).exclude(status=AuditStatus.CANCELLED).select_related("workshop")
    d1, d2 = request.GET.get("from", ""), request.GET.get("to", "")
    if d1:
        audits = audits.filter(audit_date__gte=d1)
    if d2:
        audits = audits.filter(audit_date__lte=d2)
    return render(request, "pjo/reports.html", {"audits": audits[:300], "workshops": wss, "code": code, "d1": d1, "d2": d2})


# ------------------------------------------------------------------ print / pdf
def print_context(a, embed=False):
    snap = a.snapshot
    marks = {(m.item_key, m.occurrence): m.value for m in a.marks.all()}
    n = snap["s1"].get("frequency_count") or 12
    occs = list(range(1, n + 1))
    timing = {}
    for t in a.timings.filter(saved=True, discarded=False):
        timing[t.observation] = S.timing_summary(t)
    s1_rows = []
    seen_minor = False
    for it in snap["s1"]["items"]:
        first_minor = it["severity"] == "minor" and not seen_minor
        seen_minor = seen_minor or it["severity"] == "minor"
        cells = []
        for o in occs:
            v = SYM[marks.get((it["key"], o), "")]
            if it["kind"] == "timing" and o in timing and timing[o]["avg"] is not None:
                v = (v + " " if v else "") + f"{timing[o]['avg']:.1f}s"
            cells.append(v)
        s1_rows.append({**it, "cells": cells, "first_minor": first_minor})
    s2 = []
    for sec in snap["s2"]["sections"]:
        s2.append({**sec, "items": [{**i, "mark": SYM[marks.get((i["key"], 0), "")]} for i in sec["items"]]})
    s2_issues = list(a.issues.filter(sheet="s2").order_by("id")) + [o.issue for o in a.issue_occurrences.filter(sheet="s2").exclude(issue__audit=a)]
    ans = {x.item_key: x for x in a.s3_answers.all()}
    s3 = [{**q, "a": ans.get(q["key"])} for q in snap["s3"]["questions"]]
    reviews = {r.case_key: r for r in a.s4_reviews.exclude(case_key="")}
    s4 = [{**c, "r": reviews.get(c["key"])} for c in snap["s4"]["cases"]]
    own = a.s4_reviews.filter(case_key="")
    sigs = {}
    for s in S.active_signatures(a).filter(followup__isnull=True):
        img = ""
        if embed:
            p = Path(settings.PRIVATE_MEDIA_ROOT) / s.image_path
            if p.exists():
                img = "data:image/png;base64," + base64.b64encode(p.read_bytes()).decode()
        else:
            img = f"/files/signature/{s.id}/"
        sigs[s.role] = {"s": s, "img": img}
    s1_issues = a.issue_occurrences.filter(sheet="s1").select_related("issue").order_by("id")
    followups = a.major_followups.all()
    return {
        "a": a, "snap": snap, "occs": occs, "s1_rows": s1_rows, "timing": sorted(timing.values(), key=lambda x: x["observation"]),
        "s2": s2, "s2_issues": s2_issues, "s3": s3, "suggestions": a.suggestions.all(), "s4": s4, "own": own, "sigs": sigs,
        "s1_issues": s1_issues, "followups": followups, "embed": embed, "font_dir": FONT_DIR.as_uri(),
        "printed_at": timezone.localtime(), "gbr": S.gbr_mapping(), "revisions": a.revisions.all(),
    }


@login_required
def audit_print(request, pk):
    a = get_object_or_404(Audit, pk=pk)
    perm.require(perm.can_view_audit(request.user, a))
    return render(request, "pjo/print_audit.html", print_context(a))


def render_pdf(html):
    from playwright.sync_api import sync_playwright

    kw = {"args": ["--no-sandbox"]}
    exe = settings.CHROMIUM_EXECUTABLE
    if exe and os.path.exists(exe):
        kw["executable_path"] = exe
    with sync_playwright() as p:
        b = p.chromium.launch(**kw)
        page = b.new_page()
        page.set_content(html, wait_until="load")
        page.evaluate("document.fonts.ready")
        pdf = page.pdf(format="A4", landscape=True, print_background=True,
                       margin={"top": "8mm", "bottom": "10mm", "left": "8mm", "right": "8mm"},
                       display_header_footer=True, header_template="<span></span>",
                       footer_template='<div style="font-size:8px;width:100%;text-align:right;padding-right:10mm">'
                                       '<span class="pageNumber"></span>/<span class="totalPages"></span></div>')
        b.close()
    return pdf


@login_required
def audit_pdf(request, pk):
    a = get_object_or_404(Audit, pk=pk)
    perm.require(perm.can_view_audit(request.user, a))
    html = render_to_string("pjo/print_audit.html", print_context(a, embed=True), request=request)
    try:
        pdf = render_pdf(html)
    except Exception as e:  # Chromium missing in this deployment
        log.exception("PDF render failed")
        return HttpResponse(f"สร้าง PDF บนเซิร์ฟเวอร์ไม่ได้ ({e.__class__.__name__}) กรุณาเปิดหน้าพิมพ์แล้วเลือก Save as PDF: "
                            f"/reports/audits/{a.pk}/print/", status=503, content_type="text/plain; charset=utf-8")
    S.trail(request.user, "export.pdf", audit=a)
    resp = HttpResponse(pdf, content_type="application/pdf")
    resp["Content-Disposition"] = f'attachment; filename="{a.audit_no}_rev{a.revision}.pdf"'
    return resp


# ------------------------------------------------------------------ excel
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill("solid", fgColor="DCE6F1")
WRAP = Alignment(wrap_text=True, vertical="top")


def _sheet(wb, title, header, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill, c.border, c.alignment = Font(bold=True), HEAD, BORDER, WRAP
    for r in rows:
        ws.append(["" if v is None else v for v in r])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border, c.alignment = BORDER, WRAP
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


def _dt(v):
    return timezone.localtime(v).strftime("%Y-%m-%d %H:%M") if v else ""


def issue_rows(qs):
    return [[i.issue_no, i.audit.audit_no if i.audit else "", i.workshop.code, str(i.team or ""), i.station, i.model,
             str(i.employee or ""), i.sheet, i.item_code, ", ".join(str(o.occurrence) for o in i.occurrences.all() if o.occurrence),
             i.get_issue_type_display(), i.priority, i.title, i.description, i.attachments.count(), i.owner_name or str(i.owner or ""),
             i.due_date.isoformat() if i.due_date else "", i.containment, i.root_cause, i.permanent_action, i.effectiveness,
             str(i.verifier or ""), _dt(i.closed_at), i.status_th, "เกินกำหนด" if i.is_overdue else "", i.occurrences.count(),
             i.gbr] for i in qs]


ISSUE_HEADER = ["Issue ID", "Audit ID", "Workshop", "Team", "Station", "Model", "พนักงาน", "ระยะ/ตาราง", "รหัสข้อ", "ครั้งที่พบ",
                "ประเภทปัญหา", "ความสำคัญ", "หัวข้อ", "รายละเอียด", "จำนวนหลักฐาน", "ผู้รับผิดชอบ", "Due Date", "การแก้ไขทันที",
                "สาเหตุ", "การแก้ไขถาวร", "ผลตรวจประสิทธิผล", "ผู้ยืนยัน", "วันปิด", "สถานะ", "Overdue", "จำนวนครั้งที่พบ", "G/B/R"]
ISSUE_WIDTHS = [16, 18, 10, 14, 10, 10, 18, 8, 8, 8, 18, 8, 30, 40, 8, 16, 11, 30, 30, 30, 30, 14, 16, 14, 10, 8, 6]


@login_required
def audit_xlsx(request, pk):
    a = get_object_or_404(Audit, pk=pk)
    perm.require(perm.can_view_audit(request.user, a))
    wb = Workbook()
    wb.remove(wb.active)
    snap = a.snapshot
    info = [("เลข Audit", a.audit_no), ("Revision", a.revision), ("สถานะ", a.status_th), ("วันที่", a.audit_date.isoformat()),
            ("กะ", a.shift), ("Workshop", a.workshop.name), ("Group", a.group), ("Team", str(a.team or "")), ("Line", a.line),
            ("Station", a.station), ("Model", a.model), ("พนักงานที่ถูกตรวจ", a.employee_name), ("ผู้ตรวจ", str(a.auditor)),
            ("TL/GL", a.tlgl_name), ("แม่แบบมาตรฐาน", a.template.label), ("ไฟล์ต้นฉบับ", snap.get("source_file", "")),
            ("OIS/WES Revision", a.ois_revision), ("OIS เป็นฉบับปัจจุบัน", {"yes": "Yes", "no": "No"}.get(a.ois_current, "ยังไม่ตอบ")),
            ("Standard Time (s)", a.standard_time_s), ("แหล่งที่มา Standard Time", a.standard_time_source),
            ("Takt Time (s)", a.takt_time_s), ("แหล่งที่มา Takt Time", a.takt_time_source), ("ส่งออกเมื่อ", _dt(timezone.now()))]
    _sheet(wb, "ข้อมูล Audit", ["หัวข้อ", "ค่า"], info, [26, 60])
    n = snap["s1"].get("frequency_count") or 12
    marks = {(m.item_key, m.occurrence): m for m in a.marks.select_related("recorded_by")}
    rows = []
    for it in snap["s1"]["items"]:
        rows.append([it["no"], "หลัก" if it["severity"] == "major" else "รอง", it["text"]] +
                    [SYM[marks[(it["key"], o)].value] if (it["key"], o) in marks else "" for o in range(1, n + 1)])
    _sheet(wb, "ตาราง1 สังเกตงาน", ["ข้อ", "ประเภท", "รายการ"] + [f"ครั้งที่ {o}" for o in range(1, n + 1)], rows, [5, 7, 70] + [7] * n)
    trows = []
    for t in a.timings.filter(saved=True).prefetch_related("laps"):
        for l in t.laps.all():
            trows.append([t.observation, l.lap_no, float(l.seconds), float(l.gross_seconds), float(l.paused_seconds), "; ".join(l.pause_reasons),
                          "Manual" if l.source == "manual" else "Stopwatch", l.manual_reason, "ใช้" if l.valid else "ไม่ใช้",
                          l.invalid_reason, "ถูกแทนที่" if t.discarded else "", _dt(t.saved_at), str(t.saved_by or ""), "; ".join(w.get("message", "") for w in t.warnings)])
    _sheet(wb, "การจับเวลา", ["ครั้งที่สังเกต", "รอบจับเวลา", "เวลาสุทธิ (s)", "เวลารวม (s)", "หยุดชั่วคราว (s)", "เหตุผลหยุด",
                              "ที่มา", "เหตุผล Manual", "สถานะรอบ", "เหตุผลไม่ใช้", "หมายเหตุ", "บันทึกเมื่อ", "ผู้บันทึก", "คำเตือน"], trows,
           [10, 10, 12, 12, 12, 20, 11, 20, 9, 20, 12, 16, 14, 30])
    rows = []
    for sec in snap["s2"]["sections"]:
        for i in sec["items"]:
            m = marks.get((i["key"], 0))
            rows.append([sec["title"], i["code"], i["text"], SYM[m.value] if m else ""])
    _sheet(wb, "ตาราง2 โอกาสปรับปรุง", ["หมวด", "รหัส", "รายการ", "ผล"], rows, [28, 8, 90, 6])
    ans = {x.item_key: x for x in a.s3_answers.all()}
    rows = [[q["no"], q["text"], q["expected"], getattr(ans.get(q["key"]), "employee_answer", ""), SYM[getattr(ans.get(q["key"]), "before", "")],
             getattr(ans.get(q["key"]), "coaching", ""), SYM[getattr(ans.get(q["key"]), "after", "")]] for q in snap["s3"]["questions"]]
    _sheet(wb, "ตาราง3 ความเข้าใจ", ["ข้อ", "คำถาม", "ระดับความเข้าใจที่คาดหวัง", "คำตอบพนักงาน", "ก่อนสอน", "สิ่งที่ชี้แนะ", "หลังสอน"], rows,
           [5, 45, 45, 35, 8, 35, 8])
    cases = {c["key"]: c for c in snap["s4"]["cases"]}
    rows = [[cases.get(r.case_key, {}).get("category", "ประสบการณ์ของพนักงาน"), cases.get(r.case_key, {}).get("case", r.case_text),
             cases.get(r.case_key, {}).get("answer", ""), "เกิดจริง" if r.is_real_event else "สถานการณ์สมมติ", r.employee_answer, r.coaching,
             SYM[r.result], r.improvement, r.issue.issue_no if r.issue else ""] for r in a.s4_reviews.all()]
    _sheet(wb, "ตาราง4 ความผิดปกติ", ["แนวทาง", "กรณี", "คำตอบที่พึงประสงค์", "ประเภท", "คำตอบจริง", "การชี้แนะ", "ผล", "ข้อเสนอแนะ", "Issue"], rows,
           [22, 35, 40, 12, 35, 30, 6, 30, 16])
    rows = [[s.get_role_display(), s.signer_name, s.signer_code, s.identity_method, str(s.witness or ""), _dt(s.signed_at), s.audit_revision,
             s.data_hash, "ถูกแทนที่: " + s.superseded_reason if s.superseded_at else "ใช้งาน"] for s in a.signatures.all()]
    _sheet(wb, "ลายเซ็น", ["บทบาท", "ผู้ลงนาม", "รหัส", "วิธีระบุตัวตน", "พยาน", "วันเวลา", "Revision", "Hash ข้อมูลที่เซ็น", "สถานะ"], rows,
           [30, 20, 10, 22, 14, 16, 8, 66, 30])
    _sheet(wb, "ปัญหา", ISSUE_HEADER, issue_rows(Issue.objects.filter(Q(audit=a) | Q(occurrences__audit=a)).distinct()), ISSUE_WIDTHS)
    buf = io.BytesIO()
    wb.save(buf)
    S.trail(request.user, "export.xlsx", audit=a)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{a.audit_no}_rev{a.revision}.xlsx"'
    return resp


@login_required
def issues_xlsx(request):
    wss = perm.visible_workshops(request.user)
    code = request.GET.get("workshop", "")
    scope = wss.filter(code=code) if code else wss
    qs = Issue.objects.filter(Q(workshop__in=scope) | Q(owner=request.user)).distinct().select_related("audit", "workshop", "team", "employee", "owner", "verifier")
    st = request.GET.get("status")
    if st:
        qs = qs.filter(status=st)
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "ทะเบียนปัญหา PJO", ISSUE_HEADER, issue_rows(qs), ISSUE_WIDTHS)
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="PJO_issues_{timezone.localdate():%Y%m%d}.xlsx"'
    return resp
