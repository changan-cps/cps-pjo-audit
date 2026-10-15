import base64
import io
import json
import shutil
import subprocess
import tempfile
from datetime import timedelta
from pathlib import Path

from django.conf import settings
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase, override_settings
from django.utils import timezone

from org.models import Employee, RoleAssignment, Team, Workshop
from pjo import services as S
from pjo.models import (Attachment, Audit, Issue, IssueOccurrence, Mark, MajorFollowUp, PJOTemplate, Signature,
                        TimingSession)

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "source_docs"
EXPECTED_S3 = {"GA": 9, "GA_Inspection": 9, "BS": 10, "PS": 10, "LOG": 11, "PT_Sub_Line_Engine": 9, "PT_Battery": 9}

# a tiny but non-empty PNG for signatures
_buf = io.BytesIO()
try:
    from PIL import Image, ImageDraw
    im = Image.new("RGB", (200, 80), "white")
    ImageDraw.Draw(im).line([(10, 60), (190, 20)], fill="black", width=3)
    im.save(_buf, "PNG")
except ImportError:  # pragma: no cover
    _buf.write(b"\x89PNG" + b"0" * 400)
SIG = "data:image/png;base64," + base64.b64encode(_buf.getvalue()).decode()

TMP_MEDIA = tempfile.mkdtemp(prefix="pjo-test-media-")


@override_settings(PRIVATE_MEDIA_ROOT=TMP_MEDIA, ALLOWED_HOSTS=["testserver", "localhost"])
class Base(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("setup_pjo", stdout=io.StringIO())
        cls.ga = Workshop.objects.get(code="GA")
        cls.bs = Workshop.objects.get(code="BS")
        cls.team = Team.objects.create(workshop=cls.ga, group="G1", name="T1", line="L1")
        cls.emp = Employee.objects.create(code="E1", name="พนักงาน หนึ่ง", workshop=cls.ga, team=cls.team)
        cls.emp2 = Employee.objects.create(code="E2", name="พนักงาน สอง", workshop=cls.ga, team=cls.team)

        def user(name, role, ws):
            u = User.objects.create_user(name, password="pass12345", first_name=name)
            if role:
                RoleAssignment.objects.create(user=u, role=role, workshop=ws)
            return u
        cls.auditor = user("aud", "auditor", cls.ga)
        cls.auditor2 = user("aud2", "auditor", cls.ga)
        cls.verifier = user("ver", "verifier", cls.ga)
        cls.owner = user("own", "issue_owner", cls.ga)
        cls.bs_viewer = user("bsv", "viewer", cls.bs)
        cls.admin = User.objects.create_superuser("adm", "", "pass12345")
        cls.tpl = PJOTemplate.objects.get(workshop=cls.ga, status="published")

    def setUp(self):
        self.audit = S.create_audit(self.auditor, self.tpl, {"audit_date": timezone.localdate(), "station": "ST1", "model": "M1",
                                                             "employee": self.emp, "team": self.team})

    def client_for(self, u):
        from django.test import Client
        c = Client()
        c.force_login(u)
        return c

    def post(self, c, url, data):
        return c.post(url, json.dumps(data), content_type="application/json")

    def op(self, user, **op):
        op.setdefault("op_id", S.uuid.uuid4().hex)
        return S.apply_ops(self.audit, user, [op])[0]


class TemplateFidelityTests(Base):
    def test_seven_templates_match_source_files_exactly(self):
        import sys
        sys.path.insert(0, str(ROOT / "tools"))
        from extract_templates import extract
        for f in sorted(SRC.glob("*.xlsx")):
            fresh = extract(f)
            t = PJOTemplate.objects.get(workshop__code=fresh["code"], version=1)
            self.assertEqual(t.content, fresh, f"{fresh['code']} differs from its source form")
            # published = source + only the owner-confirmed corrections
            from pjo.management.commands.setup_pjo import apply_correction
            expected = fresh
            for corr in json.loads((ROOT / "pjo/seed/corrections.json").read_text(encoding="utf8")):
                if corr["workshop"] == fresh["code"]:
                    expected = apply_correction(expected, corr)
            pub = PJOTemplate.objects.get(workshop__code=fresh["code"], status="published")
            self.assertEqual(pub.content, expected, fresh["code"])

    def test_phase2_question_counts(self):
        for code, n in EXPECTED_S3.items():
            t = PJOTemplate.objects.get(workshop__code=code, status="published")
            self.assertEqual(len(t.content["s3"]["questions"]), n, code)

    def test_workshops_use_their_own_questions(self):
        texts = {t.workshop.code: [i["text"] for i in t.content["s1"]["items"]] for t in PJOTemplate.objects.filter(version=1)}
        self.assertIn("PMHV", " ".join(texts["LOG"]))
        self.assertIn("สะเก็ดเชื่อม", " ".join(texts["BS"]))
        self.assertIn("ปืนพ่นสี", " ".join(texts["PS"]))
        self.assertIn("ไฟฟ้าสถิต", " ".join(texts["PT_Battery"]))
        self.assertEqual(len(texts["LOG"]), 9)
        gai = PJOTemplate.objects.get(workshop__code="GA_Inspection", version=1).content
        self.assertIn("ขนาดตัวอย่าง", " ".join(q["text"] for q in gai["s3"]["questions"]))

    def test_sample_marks_are_not_imported_as_results(self):
        ga = PJOTemplate.objects.get(workshop__code="GA", version=1).content
        self.assertEqual(len(ga["ignored_sample_values"]), 7)
        self.assertEqual(self.audit.marks.count(), 0)
        self.assertEqual(self.audit.issues.count(), 0)
        self.assertEqual(self.audit.s3_answers.count(), 0)

    def test_pt_battery_header_corrected_as_new_version(self):
        from pjo.models import PendingItem
        self.assertEqual(PendingItem.objects.get(code="P01-PT-BATTERY-TITLE").status, "resolved")
        v1 = PJOTemplate.objects.get(workshop__code="PT_Battery", version=1)
        v2 = PJOTemplate.objects.get(workshop__code="PT_Battery", status="published")
        self.assertEqual((v1.status, v2.version), ("retired", 3))
        q = {x["key"]: x for x in v2.content["s3"]["questions"]}
        self.assertTrue(q["S3-2"]["expected"].endswith("หรืออาจส่งผลต่อการผลิตได้"))
        self.assertEqual(q["S3-4"]["expected"], "พนักงานเข้าใจระบบ 4M Change Management")
        self.assertIn("PT Sub Line", v1.content["s1"]["title"])  # source kept as imported
        for s in ("s1", "s2", "s3", "s4"):
            self.assertIn("(PT Battery)", v2.content[s]["title"])
            self.assertNotIn("PT Sub Line", v2.content[s]["title"])
        self.assertEqual([x["text"] for x in v1.content["s3"]["questions"]], [x["text"] for x in v2.content["s3"]["questions"]])
        # PT Sub Line Engine keeps its own title
        self.assertIn("(PT Sub Line)", PJOTemplate.objects.get(workshop__code="PT_Sub_Line_Engine", status="published").content["s1"]["title"])

    def test_gbr_confirmed_and_derived_from_status(self):
        self.assertEqual({k: S.gbr_mapping()[k] for k in "GBR"}, {"G": "Complete", "B": "On progress", "R": "Overdue"})
        i, _ = S.record_finding(self.audit, self.auditor, sheet="s2", item_key="S2-A-1", occurrence=0, title="x",
                                due_date=timezone.localdate() + timedelta(days=3))
        self.assertEqual(i.gbr_auto, "B")
        i.due_date = timezone.localdate() - timedelta(days=1)
        self.assertEqual(i.gbr_auto, "R")
        i.status = "closed"
        self.assertEqual(i.gbr_auto, "G")

    def test_confirmed_items_resolved_and_open_items_remain(self):
        from pjo.models import PendingItem
        open_codes = set(PendingItem.objects.filter(status="pending").values_list("code", flat=True))
        self.assertEqual(open_codes, set())

    def test_setup_is_idempotent(self):
        call_command("setup_pjo", stdout=io.StringIO())
        call_command("setup_pjo", stdout=io.StringIO())
        self.assertEqual(PJOTemplate.objects.count(), 9)  # 7 sources + 2 confirmed PT Battery corrections


class MarkAndIssueTests(Base):
    def test_blank_means_not_checked_and_na_rules(self):
        self.assertEqual(Mark.objects.filter(audit=self.audit).count(), 0)
        r = self.op(self.auditor, type="mark", item_key="S2-A-1", occurrence=0, value="na")
        self.assertEqual(r["status"], "error")  # sheet 2 has no N/A
        r = self.op(self.auditor, type="mark", item_key="S1-1", occurrence=13, value="pass")
        self.assertEqual(r["status"], "error")

    def test_fail_creates_issue_on_right_item_without_duplicates(self):
        r = self.op(self.auditor, type="mark", item_key="S1-2", occurrence=3, value="fail", prev="")
        self.assertTrue(r["needs_issue"])
        i1, created = S.record_finding(self.audit, self.auditor, sheet="s1", item_key="S1-2", occurrence=3, title="ไม่สวม PPE", client_id="c1")
        self.assertTrue(created)
        self.assertEqual((i1.item_key, i1.issue_type), ("S1-2", "major"))
        self.assertEqual(i1.occurrences.get().occurrence, 3)
        # double tap / re-sync with same client id, and same origin with a new client id
        i2, c2 = S.record_finding(self.audit, self.auditor, sheet="s1", item_key="S1-2", occurrence=3, title="x", client_id="c1")
        i3, c3 = S.record_finding(self.audit, self.auditor, sheet="s1", item_key="S1-2", occurrence=3, title="x", client_id="c2")
        self.assertEqual((i2.id, i3.id, c2, c3), (i1.id, i1.id, False, False))
        self.assertEqual(Issue.objects.count(), 1)
        # mark again with same op id does nothing new
        op = {"type": "mark", "item_key": "S1-2", "occurrence": 3, "value": "fail", "op_id": "same"}
        S.apply_ops(self.audit, self.auditor, [op])
        res = S.apply_ops(self.audit, self.auditor, [op])[0]
        self.assertTrue(res.get("duplicate"))
        self.assertFalse(self.op(self.auditor, type="mark", item_key="S1-2", occurrence=3, value="fail")["needs_issue"])
        # major creates a follow-up
        self.assertTrue(MajorFollowUp.objects.filter(issue=i1).exists())

    def test_changing_fail_to_pass_keeps_issue_history(self):
        self.op(self.auditor, type="mark", item_key="S1-6", occurrence=1, value="fail")
        i, _ = S.record_finding(self.audit, self.auditor, sheet="s1", item_key="S1-6", occurrence=1, title="ไม่มีสมาธิ")
        self.op(self.auditor, type="mark", item_key="S1-6", occurrence=1, value="pass")
        i.refresh_from_db()
        o = i.occurrences.get()
        self.assertIn("× เป็น pass", o.mark_changed_note)
        self.assertEqual(Issue.objects.filter(pk=i.pk).count(), 1)
        self.assertTrue(i.events.filter(action="mark_changed").exists())

    def test_link_to_existing_issue_keeps_each_finding(self):
        i, _ = S.record_finding(self.audit, self.auditor, sheet="s1", item_key="S1-6", occurrence=1, title="ปัญหาเดิม")
        other = S.create_audit(self.auditor, self.tpl, {"station": "ST1"})
        j, created = S.record_finding(other, self.auditor, sheet="s1", item_key="S1-6", occurrence=4, title="", link_issue=i)
        self.assertFalse(created)
        self.assertEqual(j.id, i.id)
        self.assertEqual(i.occurrences.count(), 2)
        self.assertTrue(i.is_repeat)

    def test_concurrent_edit_is_detected(self):
        self.op(self.auditor, type="mark", item_key="S1-1", occurrence=1, value="pass", prev="")
        r = self.op(self.auditor2, type="mark", item_key="S1-1", occurrence=1, value="fail", prev="")
        self.assertEqual(r["status"], "conflict")
        self.assertEqual(r["current"], "pass")
        lv = Audit.objects.get(pk=self.audit.pk).lock_version
        self.assertEqual(self.op(self.auditor, type="header", fields={"station": "A"}, lock_version=lv)["status"], "ok")
        self.assertEqual(self.op(self.auditor2, type="header", fields={"station": "B"}, lock_version=lv)["status"], "conflict")

    def test_suggestion_is_separate_from_abnormality(self):
        self.op(self.auditor, type="suggestion", client_id="s-1", fields={"description": "ชั้นวางไกล", "suggestion": "ย้ายใกล้"})
        self.assertEqual(self.audit.suggestions.count(), 1)
        self.assertEqual(Issue.objects.count(), 0)

    def test_s4_example_cases_never_create_issues_automatically(self):
        self.op(self.auditor, type="s4", client_id="r1", fields={"case_key": "S4-1", "employee_answer": "แจ้งหัวหน้า", "result": "fail"})
        self.assertEqual(Issue.objects.count(), 0)
        self.op(self.auditor, type="s4", client_id="r2", fields={"case_key": "", "case_text": "เคยเกิดจริง", "is_real_event": True})
        c = self.client_for(self.auditor)
        r = self.post(c, f"/api/audits/{self.audit.pk}/finding/", {"sheet": "s4", "review_id": "r2", "title": "ชิ้นงานมีรอย", "client_id": "f-r2"})
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(Issue.objects.get().issue_type, "real_event")
        self.assertEqual(self.audit.s4_reviews.get(client_id="r2").issue_id, Issue.objects.get().id)


class IssueWorkflowTests(Base):
    def setUp(self):
        super().setUp()
        self.issue, _ = S.record_finding(self.audit, self.auditor, sheet="s2", item_key="S2-A-1", occurrence=0, title="แสงสว่างไม่พอ",
                                         owner=self.owner, due_date=timezone.localdate() - timedelta(days=1))

    def test_full_workflow_and_verification(self):
        i = self.issue
        self.assertEqual(i.status, "assigned")
        self.assertTrue(i.is_overdue)  # overdue is computed, not a workflow state
        with self.assertRaises(S.RuleError):
            S.transition_issue(i, self.owner, "closed")
        S.transition_issue(i, self.owner, "in_progress")
        with self.assertRaises(S.RuleError):
            S.transition_issue(i, self.owner, "pending_verification")  # no fix recorded
        i.permanent_action = "เพิ่มโคมไฟ"
        i.save()
        S.transition_issue(i, self.owner, "pending_verification")
        # owner cannot verify through the web
        c = self.client_for(self.owner)
        r = c.post(f"/issues/{i.pk}/", {"action": "transition", "to": "closed", "lock_version": i.lock_version})
        i.refresh_from_db()
        self.assertEqual(i.status, "pending_verification")
        c = self.client_for(self.verifier)
        r = c.post(f"/issues/{i.pk}/", {"action": "transition", "to": "closed", "lock_version": i.lock_version, "effectiveness": "วัดแสงได้ 500 lux"})
        i.refresh_from_db()
        self.assertEqual(i.status, "closed")
        self.assertEqual(i.verifier, self.verifier)
        with self.assertRaises(S.RuleError):
            S.transition_issue(i, self.verifier, "reopened", "")
        S.transition_issue(i, self.verifier, "reopened", "พบซ้ำ")
        self.assertEqual(i.status, "reopened")
        self.assertGreaterEqual(i.events.count(), 5)

    def test_issue_followup_does_not_overwrite_signed_results(self):
        self.op(self.auditor, type="mark", item_key="S1-1", occurrence=1, value="pass")
        self.audit.ois_current = "yes"
        self.audit.save()
        S.submit_for_signatures(self.audit, self.auditor)
        S.sign(self.audit, self.auditor, role="auditor", data_url=SIG, client_id="sa", signer_name="", identity_method="login", strokes=3)
        before = S.audit_data(self.audit)
        i = self.issue
        i.permanent_action = "แก้แล้ว"
        i.save()
        S.transition_issue(i, self.owner, "in_progress")
        self.assertEqual(S.audit_data(self.audit), before)


class TimingTests(Base):
    def events(self):
        t0 = 1_700_000_000_000
        return [{"type": "start", "t": t0}, {"type": "lap", "t": t0 + 30_500},
                {"type": "pause", "t": t0 + 40_000, "reason": "รอชิ้นงาน"}, {"type": "resume", "t": t0 + 45_000},
                {"type": "lap", "t": t0 + 62_250}, {"type": "stop", "t": t0 + 90_000}]

    def test_laps_from_timestamps_with_pause_deducted_and_reported(self):
        laps = S.laps_from_events(self.events())
        self.assertEqual([round(l["net"], 2) for l in laps], [30.5, 26.75, 27.75])
        self.assertEqual(laps[1]["paused"], 5.0)
        self.assertEqual(laps[1]["pause_reasons"], ["รอชิ้นงาน"])

    def test_save_with_invalid_and_manual_laps(self):
        base = {"item_key": "S1-10", "observation": 2, "client_id": "tm1", "events": self.events(), "final": True}
        with self.assertRaises(S.RuleError):
            S.save_timing(self.audit, self.auditor, {**base, "laps": [{"lap_no": 3, "valid": False}]})
        with self.assertRaises(S.RuleError):
            S.save_timing(self.audit, self.auditor, {**base, "manual_laps": [{"seconds": 25}]})
        s = S.save_timing(self.audit, self.auditor, {**base, "laps": [{"lap_no": 3, "valid": False, "invalid_reason": "ทำผิดลำดับ"}],
                                                     "manual_laps": [{"seconds": "29.40", "reason": "วัดจากวิดีโอ"}]})
        sm = S.timing_summary(s)
        self.assertEqual((sm["count_used"], sm["count_total"]), (3, 4))
        self.assertAlmostEqual(sm["avg"], (30.5 + 26.75 + 29.4) / 3, places=2)
        self.assertEqual((sm["min"], sm["max"]), (26.75, 30.5))
        self.assertEqual(sm["evaluation"], "รอประเมิน")  # no invented ±10% rule
        self.assertEqual(s.laps.get(lap_no=3).invalid_reason, "ทำผิดลำดับ")  # kept, not deleted
        self.assertEqual(s.laps.get(lap_no=4).source, "manual")

    def test_new_measurement_replaces_but_keeps_old(self):
        base = {"item_key": "S1-10", "observation": 1, "events": self.events(), "final": True}
        S.save_timing(self.audit, self.auditor, {**base, "client_id": "a"})
        S.save_timing(self.audit, self.auditor, {**base, "client_id": "b"})
        self.assertEqual(TimingSession.objects.filter(audit=self.audit, discarded=True).count(), 1)
        self.assertEqual(TimingSession.objects.filter(audit=self.audit, discarded=False, saved=True).count(), 1)

    def test_timing_only_on_timing_item(self):
        with self.assertRaises(S.RuleError):
            S.save_timing(self.audit, self.auditor, {"item_key": "S1-1", "observation": 1, "client_id": "x", "events": [], "final": True})


class SignatureTests(Base):
    def ready(self):
        self.op(self.auditor, type="mark", item_key="S1-1", occurrence=1, value="fail")
        self.issue, _ = S.record_finding(self.audit, self.auditor, sheet="s1", item_key="S1-1", occurrence=1, title="ไม่สวม PPE")
        self.audit.ois_current = "yes"
        self.audit.save()
        S.submit_for_signatures(self.audit, self.auditor)
        self.audit.refresh_from_db()

    def test_submit_requires_issue_for_every_fail(self):
        self.op(self.auditor, type="mark", item_key="S1-5", occurrence=1, value="fail")
        self.audit.ois_current = "yes"
        self.audit.save()
        with self.assertRaises(S.RuleError) as e:
            S.submit_for_signatures(self.audit, self.auditor)
        self.assertEqual(e.exception.code, "unlogged")

    def test_empty_signature_rejected_and_employee_must_match(self):
        self.ready()
        with self.assertRaises(S.RuleError):
            S.sign(self.audit, self.auditor, role="employee_major", data_url=SIG, client_id="x", signer_name="", signer_code="E1",
                   identity_method="employee_code_witnessed", strokes=0)
        with self.assertRaises(S.RuleError):
            S.sign(self.audit, self.auditor, role="employee_major", data_url=SIG, client_id="x", signer_name="", signer_code="E2",
                   identity_method="employee_code_witnessed", strokes=2)
        sig = S.sign(self.audit, self.auditor, role="employee_major", data_url=SIG, client_id="x", signer_name="", signer_code="E1",
                     identity_method="employee_code_witnessed", strokes=2)
        self.assertEqual((sig.signer_name, sig.witness, sig.identity_method), ("พนักงาน หนึ่ง", self.auditor, "employee_code_witnessed"))
        self.assertEqual(len(sig.data_hash), 64)
        self.assertEqual(sig.audit_revision, 1)

    def test_complete_lock_amend_resign(self):
        self.ready()
        for role, kw in [("employee_major", {"signer_code": "E1", "identity_method": "employee_code_witnessed"}),
                         ("employee_s3", {"signer_code": "E1", "identity_method": "employee_code_witnessed"}),
                         ("supervisor_s3", {"identity_method": "login"}), ("auditor", {"identity_method": "login"})]:
            S.sign(self.audit, self.auditor, role=role, data_url=SIG, client_id="c-" + role, signer_name="", strokes=3, **kw)
        self.audit.refresh_from_db()
        self.assertEqual(self.audit.status, "completed")
        self.assertTrue(self.audit.issues.exclude(status="closed").exists())  # completed ≠ issues closed
        r = self.op(self.auditor, type="mark", item_key="S1-1", occurrence=2, value="pass")
        self.assertEqual(r["code"], "locked")
        S.amend(self.audit, self.auditor, "บันทึกครั้งที่ 2 ตกหล่น", ["s1"])
        self.audit.refresh_from_db()
        self.assertEqual((self.audit.revision, self.audit.status), (2, "amended"))
        self.assertEqual(self.audit.revisions.get().data_before["s1"]["marks"]["S1-1"], {"1": "fail"})
        superseded = Signature.objects.filter(audit=self.audit, superseded_at__isnull=False).values_list("role", flat=True)
        self.assertEqual(set(superseded), {"employee_major", "auditor"})  # s3 signatures not affected
        self.assertEqual(Signature.objects.filter(audit=self.audit).count(), 4)  # nothing deleted
        self.assertEqual(self.op(self.auditor, type="mark", item_key="S1-1", occurrence=2, value="pass")["status"], "ok")
        self.assertEqual(self.op(self.auditor, type="s3", item_key="S3-1", field="coaching", value="x")["code"], "locked")
        S.submit_for_signatures(self.audit, self.auditor)
        S.sign(self.audit, self.auditor, role="employee_major", data_url=SIG, client_id="re1", signer_name="", signer_code="E1",
               identity_method="employee_code_witnessed", strokes=3)
        S.sign(self.audit, self.auditor, role="auditor", data_url=SIG, client_id="re2", signer_name="", identity_method="login", strokes=3)
        self.audit.refresh_from_db()
        self.assertEqual(self.audit.status, "completed")
        self.assertEqual(Signature.objects.filter(audit=self.audit, audit_revision=2).count(), 2)

    def test_signature_not_reused_for_new_audit(self):
        self.ready()
        S.sign(self.audit, self.auditor, role="auditor", data_url=SIG, client_id="only-once", signer_name="", identity_method="login", strokes=3)
        other = S.create_audit(self.auditor, self.tpl, {})
        self.assertFalse(Signature.objects.filter(audit=other).exists())
        # replaying the same client id returns the original signature, never a copy on another audit
        again = S.sign(self.audit, self.auditor, role="auditor", data_url=SIG, client_id="only-once", signer_name="", identity_method="login", strokes=3)
        self.assertEqual(again.audit_id, self.audit.id)

    def test_major_recertification_flow(self):
        self.ready()
        f = MajorFollowUp.objects.get(issue=self.issue)
        with self.assertRaises(S.RuleError):
            S.sign(self.audit, self.auditor, role="recertifier", data_url=SIG, client_id="rc0", signer_name="", identity_method="login",
                   followup=f, strokes=3)
        self.op(self.auditor, type="followup", id=f.id, fields={"supervisor_notified": "GL สมศักดิ์ 09:10", "training_detail": "สอนการสวม PPE",
                                                               "recert_by": "GL สมศักดิ์", "recert_result": "สวมถูกต้อง 3 รอบ"})
        f.refresh_from_db()
        self.assertEqual(f.stage, "trained")
        S.sign(self.audit, self.auditor, role="recertifier", data_url=SIG, client_id="rc1", signer_name="", identity_method="login", followup=f, strokes=3)
        S.sign(self.audit, self.auditor, role="employee_recert", data_url=SIG, client_id="rc2", signer_name="", signer_code="E1",
               identity_method="employee_code_witnessed", followup=f, strokes=3)
        f.refresh_from_db()
        self.assertEqual(f.stage, "recertified")
        # first acknowledgement and recert acknowledgement are separate signatures
        self.assertNotEqual(Signature.objects.get(client_id="rc2").role, "employee_major")
        self.assertEqual(self.op(self.auditor, type="followup", id=f.id, fields={"recert_result": "x"})["status"], "error")


class PermissionTests(Base):
    def test_other_workshop_user_cannot_read_or_write(self):
        c = self.client_for(self.bs_viewer)
        self.assertEqual(c.get(f"/audits/{self.audit.pk}/").status_code, 403)
        self.assertEqual(c.get(f"/api/audits/{self.audit.pk}/state/").status_code, 403)
        self.assertEqual(self.post(c, f"/api/audits/{self.audit.pk}/ops/", {"ops": []}).status_code, 403)
        self.assertEqual(c.get(f"/reports/audits/{self.audit.pk}/xlsx/").status_code, 403)
        self.assertEqual(c.get(f"/reports/audits/{self.audit.pk}/print/").status_code, 403)
        r = c.get("/reports/issues.xlsx")
        self.assertEqual(r.status_code, 200)
        from openpyxl import load_workbook
        self.assertEqual(load_workbook(io.BytesIO(r.content)).active.max_row, 1)  # header only

    def test_private_files_require_permission(self):
        i, _ = S.record_finding(self.audit, self.auditor, sheet="s2", item_key="S2-A-1", occurrence=0, title="x")
        att = S.save_attachment(self.auditor, SimpleUploadedFile("a.png", _buf.getvalue(), content_type="image/png"), issue=i)
        self.assertNotIn(str(Path(settings.PRIVATE_MEDIA_ROOT)), settings.STATIC_ROOT.as_posix())
        self.assertEqual(self.client_for(self.auditor).get(f"/files/attachment/{att.id}/").status_code, 200)
        self.assertEqual(self.client_for(self.bs_viewer).get(f"/files/attachment/{att.id}/").status_code, 403)
        from django.test import Client
        self.assertEqual(Client().get(f"/files/attachment/{att.id}/").status_code, 302)  # login required
        with self.assertRaises(S.RuleError):
            S.save_attachment(self.auditor, SimpleUploadedFile("x.exe", b"MZ" * 100, content_type="application/x-msdownload"), issue=i)

    def test_viewer_and_owner_cannot_edit_audit(self):
        for u in (self.owner, self.verifier):
            r = self.post(self.client_for(u), f"/api/audits/{self.audit.pk}/ops/", {"ops": [{"type": "mark", "item_key": "S1-1", "occurrence": 1, "value": "pass"}]})
            self.assertEqual(r.status_code, 403)

    def test_only_admin_manages_templates(self):
        r = self.client_for(self.auditor).post(f"/standards/{self.tpl.pk}/", {"action": "new_draft"})
        self.assertEqual(r.status_code, 403)

    def test_employee_login_signature_only_for_audited_employee(self):
        u = User.objects.create_user("e2user", password="pass12345")
        self.emp2.user = u
        self.emp2.save()
        self.op(self.auditor, type="mark", item_key="S1-1", occurrence=1, value="pass")
        self.audit.ois_current = "yes"
        self.audit.save()
        S.submit_for_signatures(self.audit, self.auditor)
        r = self.post(self.client_for(u), f"/api/audits/{self.audit.pk}/sign/", {"role": "employee_s3", "identity_method": "login",
                                                                                 "image": SIG, "strokes": 3, "client_id": "z"})
        self.assertEqual(r.status_code, 403)


class TemplateVersionTests(Base):
    def test_new_version_does_not_change_old_audits(self):
        old_snapshot = json.loads(json.dumps(self.audit.snapshot))
        c = self.client_for(self.admin)
        c.post(f"/standards/{self.tpl.pk}/", {"action": "new_draft"})
        draft = PJOTemplate.objects.get(workshop=self.ga, status="draft")
        c.post(f"/standards/{draft.pk}/", {"action": "edit_item", "key": "S1-1", "text": "ข้อความใหม่", "allow_na": ""})
        c.post(f"/standards/{draft.pk}/", {"action": "publish", "change_note": "แก้ถ้อยคำข้อ 1"})
        draft.refresh_from_db()
        self.tpl.refresh_from_db()
        self.assertEqual((draft.status, self.tpl.status), ("published", "retired"))
        self.audit.refresh_from_db()
        self.assertEqual(self.audit.snapshot, old_snapshot)
        self.assertNotEqual(self.audit.snapshot["s1"]["items"][0]["text"], "ข้อความใหม่")
        new = S.create_audit(self.auditor, draft, {})
        self.assertEqual(new.snapshot["s1"]["items"][0]["text"], "ข้อความใหม่")
        self.assertFalse(new.snapshot["s1"]["items"][0]["allow_na"])
        r = S.apply_ops(new, self.auditor, [{"type": "mark", "item_key": "S1-1", "occurrence": 1, "value": "na", "op_id": "na1"}])[0]
        self.assertEqual(r["status"], "error")

    def test_cannot_change_workshop_after_answers(self):
        bs_tpl = PJOTemplate.objects.get(workshop=self.bs, status="published")
        self.op(self.auditor, type="mark", item_key="S1-1", occurrence=1, value="pass")
        with self.assertRaises(S.RuleError) as e:
            S.change_template(self.auditor, self.audit, bs_tpl)
        self.assertEqual(e.exception.code, "template_locked")
        empty = S.create_audit(self.auditor, self.tpl, {})
        S.change_template(self.auditor, empty, bs_tpl)
        self.assertEqual(empty.snapshot["code"], "BS")


class ExportTests(Base):
    def setUp(self):
        super().setUp()
        self.op(self.auditor, type="mark", item_key="S1-1", occurrence=1, value="pass")
        self.op(self.auditor, type="s3", item_key="S3-1", field="employee_answer", value="ประกอบตามลำดับ")

    def test_excel_contains_thai_results_and_versions(self):
        from openpyxl import load_workbook
        r = self.client_for(self.auditor).get(f"/reports/audits/{self.audit.pk}/xlsx/")
        wb = load_workbook(io.BytesIO(r.content))
        self.assertIn("ตาราง1 สังเกตงาน", wb.sheetnames)
        info = {row[0].value: row[1].value for row in wb["ข้อมูล Audit"].iter_rows(min_row=2)}
        self.assertEqual(info["แม่แบบมาตรฐาน"], self.tpl.label)
        self.assertEqual(wb["ตาราง1 สังเกตงาน"]["D2"].value, "√")
        self.assertEqual(wb["ตาราง3 ความเข้าใจ"]["D2"].value, "ประกอบตามลำดับ")

    def test_pdf_has_four_sheets_and_readable_thai(self):
        try:
            import playwright  # noqa
        except ImportError:
            self.skipTest("Playwright not installed")
        r = self.client_for(self.auditor).get(f"/reports/audits/{self.audit.pk}/pdf/")
        if r.status_code == 503:
            self.skipTest("Chromium not available")
        self.assertEqual(r["Content-Type"], "application/pdf")
        if not shutil.which("pdftotext"):
            return
        with tempfile.NamedTemporaryFile(suffix=".pdf") as f:
            f.write(r.content)
            f.flush()
            text = subprocess.run(["pdftotext", "-layout", f.name, "-"], capture_output=True, text=True).stdout
        for needle in ("ระยะที่ 1 การสังเกตการทำงานเป็นระยะ (ส่วนที่ 1)", "ระยะที่ 1 การสังเกตการทำงานเป็นระยะ (ส่วนที่ 2)",
                       "ระยะที่ 2 ความเข้าใจเชิงลึก", "ระยะที่ 3 การจัดการกรณีความผิดปกติ", "ประกอบตามลำดับ", self.audit.audit_no):
            self.assertIn(needle, text)
