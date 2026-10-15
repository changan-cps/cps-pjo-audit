"""Extract the 7 PJO Evaluation Excel forms into structured template JSON.

Every text keeps its source reference (file / worksheet / cell) so the app can
show where a question came from. Sample marks left in the forms (X, /, A-1 ...)
are reported but never imported as results.

Usage:  python tools/extract_templates.py source_docs pjo/seed/templates
"""
import json
import re
import sys
import zipfile
from pathlib import Path

import openpyxl

WORKSHOPS = {
    "GA": "GA",
    "GA_Inspection": "GA Inspection",
    "BS": "BS",
    "PS": "PS",
    "LOG": "LOG",
    "PT_Sub_Line_Engine": "PT Sub Line Engine",
    "PT_Battery": "PT Battery",
}
SAMPLE_MARKS = {"X", "x", "/", "√", "×", "A-1"}


def clean(v):
    if v is None:
        return ""
    s = str(v).replace("​", "")
    s = re.sub(r"[ \t]+", " ", s)
    s = "\n".join(line.strip() for line in s.splitlines())
    return s.strip()


def one_line(s):
    return re.sub(r"\s+", " ", s).strip()


def leading_no(text):
    m = re.match(r"^\s*(\d+)\s*[.．]?", text)
    return int(m.group(1)) if m else None


class Src:
    def __init__(self, fname, ws):
        self.fname, self.ws = fname, ws

    def __call__(self, coord):
        return {"file": self.fname, "sheet": self.ws.title.strip(), "cell": coord}


def drawing_texts(path, sheet_index):
    """Text boxes in a sheet's DrawingML (openpyxl drops them). Ordered by row, col."""
    out = []
    with zipfile.ZipFile(path) as z:
        rels_name = f"xl/worksheets/_rels/sheet{sheet_index + 1}.xml.rels"
        if rels_name not in z.namelist():
            return out
        rels = z.read(rels_name).decode("utf8")
        for target in re.findall(r'Target="\.\./drawings/(drawing\d+\.xml)"', rels):
            xml = z.read(f"xl/drawings/{target}").decode("utf8")
            for sp in re.findall(r"<xdr:twoCellAnchor.*?</xdr:twoCellAnchor>", xml, re.S):
                m = re.search(r"<xdr:from><xdr:col>(\d+)</xdr:col>.*?<xdr:row>(\d+)</xdr:row>", sp, re.S)
                paras = []
                for p in re.findall(r"<a:p>.*?</a:p>|<a:p/>", sp, re.S):
                    t = "".join(re.findall(r"<a:t>([^<]*)</a:t>", p))
                    if t.strip():
                        paras.append(t.strip())
                text = "\n".join(paras)
                if text and m:
                    out.append({"row": int(m.group(2)) + 1, "col": int(m.group(1)) + 1, "text": clean(text)})
    out.sort(key=lambda d: (d["row"], d["col"]))
    return out


def parse_sheet1(ws, src, ignored):
    g = lambda c: clean(ws[c].value)
    items, labels = [], {}
    for col, sev in (("A", "major"), ("N", "minor")):
        for r in range(7, 16):
            t = g(f"{col}{r}")
            if not t:
                continue
            no = leading_no(t)
            if no is not None:
                kind = "timing" if "วัดเวลา" in t else "check"
                items.append({
                    "key": f"S1-{no}", "no": no, "text": t, "severity": sev, "kind": kind,
                    "allow_na": True, "source": src(f"{col}{r}"),
                })
            else:
                labels[f"{col}{r}"] = t
    items.sort(key=lambda i: i["no"])
    sig = []
    for coord, t in labels.items():
        tl = t
        if "ลายเซ็นยืนยันความผิดปกติหลักหลังการรับรองใหม่" in tl:
            role = "employee_recert"
        elif "ความผิดปกติหลัก" in tl:
            role = "employee_major"
        elif "ความผิดปกติรอง" in tl:
            role = "employee_minor"
        elif tl.startswith("ปัญหาที่พบ"):
            role = "problems"
        elif "พนักงานลงชื่อ" in tl:
            role = "employee_sign_label"
        elif "วัน/เดือน/ปี" in tl:
            role = "date_label"
        else:
            role = "other"
        sig.append({"role": role, "text": tl, "source": src(coord)})
    return {
        "sheet_name": ws.title.strip(),
        "title": g("A1"), "title_source": src("A1"),
        "header_text": g("A2"), "header_source": src("A2"),
        "legend": {"fail": g("A3"), "pass": g("C3"), "na": g("N3"), "source": src("A3")},
        "ois_check": {"text": g("A4"), "options": [g("B4"), g("F4")], "source": src("A4")},
        "major_response": {"text": g("A5"), "source": src("A5")},
        "minor_response": {"text": g("N5"), "source": src("N5")},
        "frequency_label": g("A6"),
        "frequency_count": sum(1 for c in "BCDEFGHIJKLM" if ws[f"{c}6"].value is not None),
        "items": items,
        "labels": sig,
    }


def parse_sheet2(ws, src, ignored):
    sections, current, log_title, table_header_row, note = [], None, None, None, None
    for r in range(1, ws.max_row + 1):
        b = clean(ws[f"B{r}"].value)
        k = clean(ws[f"K{r}"].value)
        for col in ("J", "R", "C"):
            v = clean(ws[f"{col}{r}"].value)
            if v in SAMPLE_MARKS and not (col == "C" and table_header_row and r == table_header_row):
                ignored.append({"value": v, "source": src(f"{col}{r}"), "reason": "เครื่องหมายตัวอย่าง/ค่าที่กรอกค้างในแบบฟอร์ม"})
        if r == 1:
            continue
        if b.startswith("หมายเหตุ"):
            note = {"text": b, "source": src(f"B{r}")}
            continue
        if table_header_row and r > table_header_row:
            continue
        m = re.match(r"^\s*([A-D])\s*[-–]\s*(.+)$", b)
        if m and not leading_no(b):
            current = {"code": m.group(1), "title": one_line(b), "source": src(f"B{r}"), "items": []}
            sections.append(current)
            continue
        if b.startswith("บันทึกโอกาส"):
            log_title = {"text": b, "source": src(f"B{r}")}
            current = None
            continue
        if b == "ลำดับ":
            table_header_row = r
            cols = []
            for c in ("B", "C", "D", "M", "N", "Q"):
                cols.append({"text": clean(ws[f"{c}{r}"].value), "source": src(f"{c}{r}")})
            continue
        for col, t in (("B", b), ("K", k)):
            if t and current is not None:
                no = leading_no(t)
                current["items"].append({
                    "no": no, "text": t, "source": src(f"{col}{r}"),
                })
    for s in sections:
        s["items"].sort(key=lambda i: (i["no"] is None, i["no"] or 0))
        seen = {}
        for i in s["items"]:
            base = f"{s['code']}-{i['no']}"
            seen[base] = seen.get(base, 0) + 1
            i["code"] = base if seen[base] == 1 else f"{base}#{seen[base]}"
            i["key"] = f"S2-{i['code']}"
    rows = 0
    if table_header_row:
        r = table_header_row + 1
        while isinstance(ws[f"B{r}"].value, (int, float)):
            rows += 1
            r += 1
    return {
        "sheet_name": ws.title.strip(),
        "title": clean(ws["B1"].value), "title_source": src("B1"),
        "sections": sections,
        "log_title": log_title,
        "table_columns": cols if table_header_row else [],
        "paper_rows": rows,
        "note": note,
    }


def parse_sheet3(ws, src, ignored):
    g = lambda c: clean(ws[c].value)
    qs, r = [], 6
    after = None
    for r in range(6, ws.max_row + 1):
        a = ws[f"A{r}"].value
        b = clean(ws[f"B{r}"].value)
        if isinstance(a, (int, float)) and g(f"C{r}") and not g(f"I{r}"):
            if b in SAMPLE_MARKS:
                ignored.append({"value": b, "source": src(f"B{r}"), "reason": "ผลตัวอย่างที่กรอกค้างในแบบฟอร์ม"})
            q = g(f"C{r}")
            qs.append({
                "key": f"S3-{int(a)}", "no": int(a), "text": q, "expected": g(f"F{r}"),
                "is_suggestion": "ข้อเสนอแนะ" in q,
                "source": src(f"C{r}"), "expected_source": src(f"F{r}"),
            })
        elif qs and after is None and g(f"A{r}"):
            after = r
            break
    sugg_title = {"text": g(f"A{after}"), "source": src(f"A{after}")}
    hdr = after + 1
    cols = [{"text": g(f"{c}{hdr}"), "source": src(f"{c}{hdr}")} for c in ("A", "C", "I", "O") if g(f"{c}{hdr}")]
    rows, r = 0, hdr + 1
    while isinstance(ws[f"A{r}"].value, (int, float)):
        rows += 1
        r += 1
    return {
        "sheet_name": ws.title.strip(),
        "title": g("A1"), "title_source": src("A1"),
        "goal": {"text": g("A2"), "source": src("A2")},
        "step": {"text": g("A3"), "source": src("A3")},
        "guidance": {"text": g("F3"), "source": src("F3")},
        "col_question": g("A5"), "col_expected": g("F5"),
        "questions": qs,
        "suggestion_title": sugg_title,
        "suggestion_columns": cols,
        "paper_rows": rows,
        "signature": {"text": g(f"A{r}"), "source": src(f"A{r}")},
    }


def parse_sheet4(ws, src, ignored, drawings):
    g = lambda c: clean(ws[c].value)
    hdr = None
    for r in range(1, ws.max_row + 1):
        if g(f"A{r}").startswith("แนวทางการคิด"):
            hdr = r
            break
    cases, cat, rec = [], None, None
    for r in range(hdr + 1, ws.max_row + 1):
        if g(f"L{r}").startswith("บันทึก"):
            rec = r
            break
        if g(f"A{r}"):
            cat = {"text": g(f"A{r}"), "source": src(f"A{r}")}
        if g(f"E{r}"):
            cases.append({
                "key": f"S4-{len(cases) + 1}", "no": len(cases) + 1,
                "category": cat["text"], "category_source": cat["source"],
                "case": g(f"E{r}"), "answer": g(f"K{r}"),
                "source": src(f"E{r}"), "answer_source": src(f"K{r}"),
            })
    steps = [{"no": i + 1, "text": g(c), "source": src(c)} for i, c in enumerate(("A5", "G5", "K5", "N5"))]
    # drawing boxes below the case table: quality duties (upper row) and safety behaviour (lower row)
    boxes = [d for d in drawings if d["row"] > rec and not re.fullmatch(r"\d", d["text"])]
    rows = sorted({d["row"] for d in boxes})
    groups = []
    for row in rows:
        groups.append([d["text"] for d in boxes if d["row"] in (row, row + 1) and d not in sum(([x for x in boxes if x["text"] in grp] for grp in groups), [])])
    groups = [grp for grp in groups if grp]
    return {
        "sheet_name": ws.title.strip(),
        "title": g("A1"), "title_source": src("A1"),
        "goal": {"text": g("A2"), "source": src("A2")},
        "lead": {"text": g("A4"), "source": src("A4")},
        "steps": steps,
        "discussion_title": {"text": g(f"A{hdr - 1}"), "source": src(f"A{hdr - 1}")},
        "columns": [g(f"A{hdr}"), g(f"E{hdr}"), g(f"K{hdr}")],
        "cases": cases,
        "record_title": {"text": g(f"L{rec}"), "source": src(f"L{rec}")},
        "record_columns": [g(f"L{rec + 1}"), g(f"P{rec + 1}")],
        "diagram_boxes": [{"text": d["text"], "source": {"file": src("A1")["file"], "sheet": ws.title.strip(), "cell": f"shape@R{d['row']}C{d['col']}"}} for d in boxes],
    }


def extract(path):
    m = re.search(r"Template_-_(.+)\.xlsx$", path.name)
    code = m.group(1)
    wb = openpyxl.load_workbook(path)
    W = wb.worksheets
    ignored = []
    s1 = parse_sheet1(W[0], Src(path.name, W[0]), ignored)
    s2 = parse_sheet2(W[1], Src(path.name, W[1]), ignored)
    s3 = parse_sheet3(W[2], Src(path.name, W[2]), ignored)
    s4 = parse_sheet4(W[3], Src(path.name, W[3]), ignored, drawing_texts(path, 3))
    return {
        "code": code,
        "name": WORKSHOPS[code],
        "form_code": "CATH-CPS009-B02",
        "source_file": path.name,
        "s1": s1, "s2": s2, "s3": s3, "s4": s4,
        "ignored_sample_values": ignored,
    }


def main(src_dir, out_dir):
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    for p in sorted(Path(src_dir).glob("*.xlsx")):
        t = extract(p)
        (out / f"{t['code']}.json").write_text(json.dumps(t, ensure_ascii=False, indent=1), encoding="utf8")
        print(t["code"], "S1", len(t["s1"]["items"]), "S2", sum(len(s["items"]) for s in t["s2"]["sections"]),
              "S3", len(t["s3"]["questions"]), "S4", len(t["s4"]["cases"]), "ignored", len(t["ignored_sample_values"]))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
